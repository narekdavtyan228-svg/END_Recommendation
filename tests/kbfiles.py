"""Builders of sample documents (PDF, DOCX, TXT) for the document base tests."""

import glob
import io

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

LAW = """Глава 1. Общие положения
1. Огневые работы проводятся не ближе 50 метров от мест хранения ЛВЖ и взрывчатых веществ.
2. Перед началом газоопасных работ выполняется анализ воздушной среды газоанализатором.
33. Работы в траншеях и шурфах выполняются после анализа воздуха на содержание вредных веществ.
1) анализ повторяют после каждого перерыва в работе;
2) результаты анализа записываются в наряд-допуск.
Глава 2. Требования к ограждению
128. Выемки ограждаются на высоту не менее одного метра либо защитным покрытием.
"""


def txt(text: str = LAW) -> bytes:
    return text.encode("utf-8")


def docx(lines: list[str] | None = None, table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for line in lines or LAW.strip().splitlines():
        document.add_paragraph(line)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                grid.cell(r, c).text = cell
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _font() -> str:
    for path in glob.glob("/usr/share/fonts/**/DejaVuSans.ttf", recursive=True):
        pdfmetrics.registerFont(TTFont("DejaVu", path))
        return "DejaVu"
    return "Helvetica"


def pdf(text: str = LAW) -> bytes:
    out = io.BytesIO()
    page = canvas.Canvas(out, pagesize=A4)
    page.setFont(_font(), 10)
    y = 800
    for line in text.strip().splitlines():
        page.drawString(40, y, line)
        y -= 14
    page.save()
    return out.getvalue()


def pdf_without_text() -> bytes:
    out = io.BytesIO()
    page = canvas.Canvas(out, pagesize=A4)
    page.rect(100, 100, 200, 200, fill=1)  # a drawing only: like a scan, no text layer
    page.save()
    return out.getvalue()
